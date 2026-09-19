#!/usr/bin/env python3
"""Create isolated, macro-free Office fixtures without starting Word or PowerPoint.

Requires python-pptx and python-docx (available in the bundled workspace Python).
Run from any directory: python tools/WindowsChecks/generate-office-fixtures.py
Existing fixture files are never overwritten. The generated manifest is an
independent expected-value oracle for native Office capture/resume checks.
"""
import argparse
import stat
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches

REPOSITORY = Path(__file__).resolve().parents[2]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--mode", choices=("editable", "read-only", "reading-restriction"), default="editable")
MODE = parser.parse_args().mode
OUTPUT = REPOSITORY / "testdata" / "generated" / {
    "editable": "office-extension",
    "read-only": "office-readonly",
    "reading-restriction": "office-restricted-readonly",
}[MODE]
PPTX = OUTPUT / "workbookmark-powerpoint-fixture.pptx"
DOCX = OUTPUT / "workbookmark-word-fixture.docx"
MANIFEST = OUTPUT / "fixture-manifest.json"


def write_new(path, contents):
    """Exclusive creation prevents replacing a document that may already be open."""
    if path.parent.resolve() != OUTPUT.resolve():
        raise ValueError("Fixture output escaped its dedicated generated directory")
    with path.open("xb") as output:
        output.write(contents)


def stable_metadata(properties):
    properties.author = "WorkBookmark synthetic test generator"
    properties.title = "WorkBookmark test fixture - no user content"
    properties.subject = "Office capture and resume integration test"
    properties.created = datetime(2026, 9, 19, tzinfo=timezone.utc)
    properties.modified = datetime(2026, 9, 19, tzinfo=timezone.utc)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    existing = [str(path) for path in (PPTX, DOCX, MANIFEST) if path.exists()]
    if existing:
        raise SystemExit("Refusing to overwrite existing fixtures: " + ", ".join(existing))

    deck = Presentation()
    deck.slide_width = Inches(10)
    deck.slide_height = Inches(5.625)
    stable_metadata(deck.core_properties)
    for number in range(1, 4):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Synthetic slide {number}"
        slide.placeholders[1].text = (
            f"WorkBookmark fixture {number}. This slide contains no user content.\n"
            "Capture slide 2, navigate to slide 1, then resume the saved bookmark."
        )
    pptx_data = BytesIO()
    deck.save(pptx_data)
    pptx_bytes = pptx_data.getvalue()
    with ZipFile(BytesIO(pptx_bytes)) as package:
        xml = ET.fromstring(package.read("ppt/presentation.xml"))
    presentation_ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    ids = [int(element.attrib["id"]) for element in xml.find(presentation_ns + "sldIdLst")]

    paragraphs = [
        "WorkBookmark synthetic Word fixture.",
        "Second paragraph: capture a caret at this paragraph's start.",
        "Third paragraph: navigate here before resuming the saved position.",
        "Final paragraph. No personal, confidential, linked, or macro content.",
    ]
    document = Document()
    stable_metadata(document.core_properties)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    for text in paragraphs:
        document.add_paragraph(text)
    if MODE == "reading-restriction":
        protection = OxmlElement("w:documentProtection")
        protection.set(qn("w:edit"), "readOnly")
        protection.set(qn("w:enforcement"), "1")
        document.settings.element.append(protection)
    docx_data = BytesIO()
    document.save(docx_data)
    docx_bytes = docx_data.getvalue()
    starts = []
    offset = 0
    for text in paragraphs:
        starts.append(offset)
        # Word main-story positions use UTF-16 characters, including each paragraph mark.
        offset += len(text.encode("utf-16-le")) // 2 + 1
    manifest = {
        "purpose": "Generated synthetic files only; no Office automation or user documents used",
        "mode": MODE,
        "powerpoint": {"file": PPTX.name, "slide_ids": ids, "capture_slide_number": 2,
                       "capture_slide_id": ids[1], "sha256": sha256(pptx_bytes).hexdigest()},
        "word": {"file": DOCX.name, "paragraph_starts": starts, "capture_word_start": starts[1],
                 "document_end": offset, "sha256": sha256(docx_bytes).hexdigest()},
        "limitations": "Generated package oracle; actual installed Office behavior still needs the native worker checks",
    }
    write_new(PPTX, pptx_bytes)
    write_new(DOCX, docx_bytes)
    write_new(MANIFEST, (json.dumps(manifest, indent=2) + "\n").encode("utf-8"))
    if MODE == "read-only":
        DOCX.chmod(stat.S_IREAD)
        PPTX.chmod(stat.S_IREAD)
    print(json.dumps({"output_directory": str(OUTPUT), **manifest}, indent=2))


if __name__ == "__main__":
    main()
